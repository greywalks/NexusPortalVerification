"""Exercise the browser's real contracts; Python is test-only, never application runtime."""
import atexit, io, json, os, re, warnings
from pathlib import Path
import requests
from openpyxl import Workbook, load_workbook
warnings.filterwarnings('ignore',message='Workbook contains no default style')

ROOT=os.getenv('USSI_NEXUS_TEST_URL',os.getenv('LOGICORE_TEST_URL','http://127.0.0.1:5000')).rstrip('/')
BASE=ROOT+'/index.cfm'
# This suite changes server state: it replaces the AMC and Philips dimension
# tables, generates invoices, files NonConforming items, imports inventory
# events that are never removed, and creates/deletes users. Never point it at an
# environment whose data matters.
if os.getenv('USSI_NEXUS_TEST_DISPOSABLE')!='1':
 raise SystemExit('Refusing to run: '+ROOT+' must be a disposable test instance. '
  'This suite replaces reference data and writes records that are not cleaned up. '
  'Set USSI_NEXUS_TEST_DISPOSABLE=1 only for a throwaway database (for example the CI runner).')
TEST_USERNAME=os.environ['USSI_NEXUS_BOOTSTRAP_USERNAME']
TEST_PASSWORD=os.environ['USSI_NEXUS_BOOTSTRAP_PASSWORD']
OUT=Path(__file__).parent/'generated'; OUT.mkdir(exist_ok=True)
s=requests.Session()
def check(r,status=200):
 assert r.status_code==status, (r.url,r.status_code,r.text[:1600])
 return r

def api(path,**kwargs):
 r=check(s.post(BASE+path,**kwargs) if kwargs else s.get(BASE+path))
 d=r.json();assert d.get('ok') is True,(path,d);return d

def book(sheets):
 w=Workbook();w.remove(w.active)
 for name,(headers,rows) in sheets.items():
  ws=w.create_sheet(name);ws.append(headers)
  for row in rows: ws.append(row)
 b=io.BytesIO();w.save(b);return b.getvalue()

def upload(path,filename,data,fields=None):
 return api(path,files={filename:(filename+'.xlsx',data)},data=fields or {})

def done(path):
 r=check(s.get(BASE+path));lines=[x[6:] for x in r.text.splitlines() if x.startswith('data: ')]
 d=json.loads(lines[-1]);assert d.get('type')=='done' and d.get('success') is True,d
 fn=d['filename'];r=check(s.get(BASE+'/download/'+fn));assert r.content[:2]==b'PK',r.text[:300];(OUT/fn).write_bytes(r.content)
 return d,load_workbook(io.BytesIO(r.content))

anonymous=requests.Session()
login=check(anonymous.get(BASE+'/login')).text
assert 'data-theme="dark"' in login and '/static/js/theme-init.js' in login
assert 'USSI Nexus' in login and '/static/ussi_global_logo.png' in login
theme_js=check(anonymous.get(ROOT+'/static/js/theme-init.js')).text
assert "localStorage.getItem('theme')" in theme_js
for asset in ['/static/ussi_nexus_mark.png','/static/ussi_global_logo.png']:
 branded=check(anonymous.get(ROOT+asset));assert branded.content.startswith(b'\x89PNG')
login_response=check(s.post(BASE+'/login',data={'username':TEST_USERNAME,'password':TEST_PASSWORD},allow_redirects=False),302)
assert login_response.headers.get('X-Content-Type-Options')=='nosniff'
assert api('/healthz')['engine']=='Lucee'
for protected in ['/services/InvoiceService.cfc','/config/serial_rules.json','/template/FedEx_Shipment_Upload_Template.xlsx','/server.json','/README.md','/BRANDING.md','/AUDIT_REPORT.md','/ENTERPRISE_READINESS.md']:check(s.get(ROOT+protected),404)
portal=check(s.get(BASE+'/')).text
assert '<script>window.LOGICORE_AUTH=' in portal
assert '"invoiceChildren"' in portal and '"isSuperadmin"' in portal
assert 'class="app-shell"' in portal and 'id="sidebar"' in portal
assert 'id="portal-nav-admin-permissions"' in portal and '>User Management</span>' in portal
assert 'id="portal-nav-my-account"' in portal and '>My Account</span>' in portal
assert 'data-theme-preference="' in portal
assert 'id="admin-permissions-link"' not in portal
assert '<title>USSI Nexus</title>' in portal and '/static/ussi_nexus_mark.png' in portal
shell_pages=[
 ('/training-tracker/','portal-nav-training-tracker'),
 ('/training-tracker/people','portal-nav-training-tracker'),
 ('/training-tracker/reports','portal-nav-training-tracker'),
 ('/training-tracker/account','portal-nav-training-tracker'),
 ('/training-tracker/admin','portal-nav-training-tracker'),
 ('/inventory-management/','portal-nav-inventory-management'),
 ('/inventory-management/import','portal-nav-inventory-management'),
 ('/inventory-management/imports','portal-nav-inventory-management'),
 ('/inventory-management/shipping/all','portal-nav-inventory-management'),
 ('/inventory-management/quality','portal-nav-inventory-management'),
 ('/admin/permissions',None),
 ('/account','portal-nav-my-account'),
]
for p,active in shell_pages:
 page=check(s.get(BASE+p)).text
 assert 'class="app-shell"' in page and 'id="sidebar"' in page,page[:1600]
 if active:assert f'id="{active}"' in page and f'id="{active}" href' in page,page[:1600]
 if p=='/admin/permissions':assert 'id="portal-nav-admin-permissions"' in page and 'side-link active' in page,page[:2000]
 assert '/static/js/sidebar-nav.js' in page and '/static/js/theme-init.js' in page,page[:1600]
 assert 'â' not in page and 'Â' not in page,page[:1600]
 if p.startswith('/training-tracker/'):
  assert 'class="sidebar-sub"' in page and 'workspace-subnav' not in page,page[:1600]

training_admin=check(s.get(BASE+'/training-tracker/admin')).text
assert 'Legacy course import' in training_admin and 'Admin only' in training_admin
template_response=check(s.get(BASE+'/training-tracker/admin/import/template'))
training_template=load_workbook(io.BytesIO(template_response.content),data_only=False)
assert training_template.sheetnames[:2]==['Courses','Example']
assert training_template['Courses']['A6'].value=='Course Date*'
assert training_template['Courses']['B6'].value=='Course Title*'
assert all(cell.value is None for row in training_template['Courses'].iter_rows(min_row=7) for cell in row)
assert anonymous.get(BASE+'/training-tracker/admin/import/template',allow_redirects=False).status_code in (302,303)

quality_page=check(s.get(BASE+'/inventory-management/quality')).text
assert '<th>Severity</th>' not in quality_page or '<th>Expected model</th>' in quality_page

# Snapshot the dimension tables this suite replaces and put them back on exit,
# including after a failed assertion, so a mistaken run does less damage.
# --- dimension snapshot begin ---
_dimension_snapshots={}
for client in ['amc','philips']:
 r=s.get(BASE+'/download_'+client+'_dimensions')
 if r.status_code==200 and r.content[:2]==b'PK':
  _dimension_snapshots[client]=(r.content,api('/get_'+client+'_dimensions').get('count'))
def _restore_dimensions():
 for client,(data,count) in _dimension_snapshots.items():
  try:
   r=s.post(BASE+'/upload_'+client+'_dimensions',files={'file':(client+'_dimensions_restore.xlsx',data)})
   restored=s.get(BASE+'/get_'+client+'_dimensions').json().get('count')
   state='restored' if r.ok and restored==count else f'RESTORE MISMATCH (HTTP {r.status_code}; {restored} models, expected {count})'
  except Exception as e:
   state=f'RESTORE FAILED: {e}'
  print(f'{client} dimensions {state}')
atexit.register(_restore_dimensions)
# --- dimension snapshot end ---

# Config round-trip tests the actual upload/download argument order and numeric data.
dims=book({'Dimensions':(['Model','Sq Footage'],[['TEST-75',20],['TEST-86',35]])})
for client in ['amc','philips']:
 upload('/upload_'+client+'_dimensions','file',dims)
 r=check(s.get(BASE+'/download_'+client+'_dimensions'))
 w=load_workbook(io.BytesIO(r.content));assert w.active['B2'].value==20
 assert api('/get_'+client+'_dimensions')['dimensions']['TEST-75']==20

# AMC: JSON manual correction must persist and trigger fresh analysis.
files={
 'receiving':('receiving.xlsx',book({'Receiving Export':(['Model','Serial Number','Received Date'],[['TEST-75','R1','08-03-2026']])})),
 'shipping':('shipping.xlsx',book({'Shipping Export':(['Model','Serial Number','Shipped Date'],[['TEST-75','S1','08-04-2026']])})),
 'inventory':('inventory.xlsx',book({'Inventory Export':(['Model','Serial Number'],[['MISSING','I1'],['MISSING','I1'],['TEST-75','I2']])}))}
d=api('/analyze_amc',files=files,data={'period_start':'2026-08-01','period_end':'2026-08-31','invoice_title':'AMC parity title','output_filename':'AMC_PARITY'})
assert d['missing_dimension_models']==['MISSING'],d
api('/confirm_amc',json={'dimensions':{'MISSING':1000}})
d,w=done('/stream_amc');assert d['filename']=='AMC_PARITY.xlsx';assert d['additional_sqft']==500,d;assert d['subtotal']==3686,d

# Philips: missing Size column is valid; manual dimensions and metadata survive confirm.
p=book({'Inventory':(['Model','Serial','Type'],[['MISSING-PH','I1','Demo'],['TEST-75','I2','Service']]),'Shipping':(['Model','Stock Level (Primary)'],[['TEST-75','Service']]),'Recieved':(['Model'],[['TEST-75']]),'Repairs':(['Model','Status'],[['75BDL','Repaired']])})
d=upload('/analyze_philips','report',p,{'output_filename':'PHILIPS_PARITY','invoice_title':'Philips parity title','parts_sqft':'100'})
assert 'MISSING-PH' in d['missing_dimension_models']
api('/confirm_philips',json={'dimensions':{'MISSING-PH':600}})
d,w=done('/stream_philips');assert d['filename']=='PHILIPS_PARITY.xlsx';assert d['demo_total_sqft']==600,d

# TCL must consume the JSON payload sent by the unchanged frontend.
p=book({'Inventory Export':(['Model','Serial Number','Grade','Rack','Bin','Received Date'],[['TV','U1','A','MAIN','X','08-01-2026'],['PART','P1','A','PARTS','X','08-02-2026']])})
d=upload('/analyze_tcl','inventory',p,{'date_from':'2026-08-01','date_to':'2026-08-31','output_filename':'TCL_PARITY'})
d_tcl_key=d['unit_groups'][0]['key'];d_tcl_box_key=d['part_groups'][0]['key']
api('/confirm_tcl',json={'unit_breakdowns':{d_tcl_key:'1'},'box_breakdowns':{d_tcl_box_key:'1'}})
d,w=done('/stream_tcl');assert d['filename']=='TCL_PARITY.xlsx';assert d['subtotal']==78.85,d

print('PASS: JSON casing, portal navigation, dimensions round-trip, AMC/Philips correction recomputation, TCL JSON confirmation and XLSX downloads')

# Worksheet contracts are part of parity, including the intentionally misspelled source name.
w=load_workbook(OUT/'AMC_PARITY.xlsx');assert w.sheetnames==['Breakdown','Received','Shipped','Excluded Items'];assert w['Breakdown']['A5'].value=='AMC parity title';assert w['Breakdown']['D8'].value=='=COUNTA(Received!C2:C1048576)';assert w['Received']['C2'].value.day==3
w=load_workbook(OUT/'PHILIPS_PARITY.xlsx');assert w['Breakdown']['D16'].value=='=COUNTA(Shipping!G2:G100000)';assert w['Shipping']['G2'].value=='TEST-75';assert w['Repairs']['G2'].value=='Yes'
w=load_workbook(OUT/'TCL_PARITY.xlsx');assert w.sheetnames==['Invoice','Line Items'];assert w['Line Items']['G2'].value==75

# FedEx preserves template columns (including duplicate MainContact), real Excel dates/time,
# truncation instead of rounding, blank-recipient defaults, and tracking identifiers.
fheaders=['Express or Ground Tracking ID','Net Charge Amount','Payor','Recipient Company','Recipient Name','Recipient Address Line 1','Recipient Address Line 2','Recipient City','Recipient State','Recipient Zip Code','Original Customer Reference','Original Ref#3/PO Number']
p=book({'Raw':(fheaders,[['001234567890123456',8.5,'','','','','','','','','PO1',''],['001234567890123456',8.5,'','','','','','','','','',''],['T2',1,'','Example','Matt Shaw','','','Melbourne','fl','32940-1234','','PO2'],['T3',0,'','','','','','','','','','']])})
d=upload('/analyze_fedex_shipment','raw',p,{'period_label':'082026','call_date':'2026-08-31','output_filename':'FEDEX_PARITY'})
api('/build_fedex_shipment',json={});d,w=done('/stream_fedex_shipment');assert d['row_count']==2,d;assert d['total_price']==11.17,d;assert d['skipped_count']==2,d
headers=[c.value for c in w.active[1]];assert len(headers)>23;assert headers.count('MainContact')==2
assert w.active.cell(2,headers.index('Summary')+1).value.endswith('001234567890123456')
assert w.active.cell(2,headers.index('CallRcvd')+1).value.day==31
assert w.active.cell(2,headers.index('CallRcvdTime')+1).value.hour==17

# Workshop: previously triaged repairs get the reduced rate; previously repaired
# serials require a later shipment; source/master companion files retain other sheets.
raw=book({'Repair Data':(['Date Integer','Actual Model','Actual Serial','Derive Size','Result','Category'],[
 ['08-05-2026','AP9-A75-NA-R','9A75XK001','75','Mainboard replaced','Refurbished'],
 ['08-05-2026','AP9-A75-NA-R','9A75XK002','75','Mainboard replaced','Refurbished'],
 ['08-06-2026','AP9-A86-NA-R','9A86XK003','86','Pending LCD','Pending Parts'],
 ['07-01-2026','AP9-A75-NA-R','9A75XK004','75','Mainboard replaced','Refurbished']]),'Notes':(['Keep'],[['untouched']])})
master=book({'Repair Log':(['Month','Model','Serial','Type'],[['2026-07-01','AP9-A75-NA-R','9A75XK002','Basic']]),'Triage Log':(['Date','Model','Serial','Type'],[['2026-07-01','AP9-A75-NA-R','9A75XK001','Triage - Basic']]),'Notes':(['Keep'],[['history']])})
ship=b'Shipped Date,Serial Number\n07-01-2026,9A75XK002\n'
d=api('/sanitize',files={'raw_file':('raw.xlsx',raw),'prev_invoiced':('master.xlsx',master),'shipping':('ship.csv',ship)},data={'date_from':'2026-08-01','date_to':'2026-08-31','invoice_date':'2026-09-01','completed_date':'2026-08-31','customer':'Promethean','call_id':'C-PARITY'})
assert d['total_records']==3,d;assert d['issue_count']==0,d
api('/generate',data={'output_filename':'WORKSHOP_PARITY','corrections':'{}'})
d,w=done('/stream');assert d['subtotal']==245,d;assert d['excluded_count']==1,d
assert w['Breakdown']['E1'].value=='C-PARITY';assert w['Depot Repair']['C2'].value=='BasicSmall - Previously Triaged';assert w['Depot Repair']['E2'].value==64
assert w['Triage Units']['D2'].value=='HeavyLarge-Triage';assert 'TblPartTesting' in w['Part Testing & Programming'].tables
for field in ['corrected_filename','master_filename']:
 data=check(s.get(BASE+'/download/'+d[field])).content;cw=load_workbook(io.BytesIO(data));assert 'Notes' in cw.sheetnames
 if field=='master_filename':assert cw['Repair Log'].max_row==3;assert cw['Triage Log'].max_row==3
 else:assert 'Sanitized Data' in cw.sheetnames

# NonConforming CRUD and export use JSON and retain punctuation in cell values.
item=api('/nonconforming/api/items',json={'model':'Parity, Model','serial':'001234','carrier':'FedEx','status':'Pending'})['item'];iid=item['id']
r=check(s.patch(BASE+f'/nonconforming/api/items/{iid}',json={'status':'Resolved','addtl_info':'First, second'}));assert r.json()['item']['status']=='Resolved',r.text
r=check(s.get(BASE+'/nonconforming/api/export',params={'q':'Parity, Model'}));w=load_workbook(io.BytesIO(r.content));assert any(row[3]=='Parity, Model' and row[4]=='001234' for row in w.active.iter_rows(min_row=2,values_only=True))
label=api(f'/nonconforming/api/items/{iid}/label');assert '^XA' in label['zpl'] and item['number'] in label['zpl']
check(s.delete(BASE+f'/nonconforming/api/items/{iid}'));check(s.get(BASE+f'/nonconforming/api/items/{iid}'),404)
print('PASS: source workbook layouts/formulas, FedEx template/date/time, Workshop dedup/rates/companions, NonConforming CRUD/export/labels')

# Training: create content and attendance, then exercise signed scans and reports.
def postform(path,data):return check(s.post(BASE+path,data=data))
r=postform('/training-tracker/week/new',{'title':'Parity Week','start_date':'2026-08-03','notes':'Test notes'});wid=re.search(r'/week/(\d+)',r.url).group(1)
postform('/training-tracker/people/add',{'name':'Parity Person','email':'parity@example.test'})
r=postform(f'/training-tracker/week/{wid}/topic/new',{'title':'Parity Topic','lesson_plan':'Test lesson','key_points':'One\nTwo'})
tid=re.findall(r'/training-tracker/topic/(\d+)/edit',r.text)[-1]
r=s.post(BASE+f'/training-tracker/topic/{tid}/video/add',data={'video_title':'Unsafe','url':'javascript:alert(1)'})
check(r,400);assert 'Video links must begin' in r.text
r=postform(f'/training-tracker/topic/{tid}/video/add',{'video_title':'Safe video','url':'https://example.test/training'})
assert 'target="_blank" rel="noopener noreferrer"' in r.text
r=postform(f'/training-tracker/week/{wid}/session/new',{'trainer_name':'Parity Trainer','session_date':'2026-08-03','location':'Workshop'})
sid=re.findall(r'href="/training-tracker/session/(\d+)"',r.text)[-1]
r=check(s.get(BASE+f'/training-tracker/session/{sid}'));pid=re.findall(r'/attendance/(\d+)"',r.text)[-1]
postform(f'/training-tracker/session/{sid}/attendance/{pid}',{'attended':'on','signature':'Parity Person'})
r=check(s.get(BASE+f'/training-tracker/session/{sid}/signoff/download'));assert r.content.startswith(b'%PDF'),r.text[:500]
pdf=r.content
postform(f'/training-tracker/session/{sid}/attendance/toggle',{'person_id':pid,'attended':'1'})
r=check(s.post(BASE+f'/training-tracker/session/{sid}/signoff/upload',files={'signoff_file':('signed.pdf',pdf,'text/html')}))
r=check(s.get(BASE+f'/training-tracker/session/{sid}/signoff/view'));assert r.content==pdf
assert r.headers.get('Content-Type','').startswith('application/pdf')
assert r.headers.get('Content-Security-Policy')=='sandbox'
r=check(s.get(BASE+'/training-tracker/reports/attendance-matrix',params={'start_week':wid,'end_week':wid}));w=load_workbook(io.BytesIO(r.content));assert any('X' in row for row in w.active.iter_rows(values_only=True))
r=check(s.get(BASE+'/training-tracker/reports/signoff-zip',params={'start_week':wid,'end_week':wid}));assert r.content.startswith(b'PK')
postform(f'/training-tracker/week/{wid}/delete',{});postform(f'/training-tracker/people/{pid}/delete',{})

# Inventory imports exercise rows with no MSO and full FedEx event details.
import uuid
serial='PARITY'+uuid.uuid4().hex[:8]
p=book({'Receiving Export':(['Received Date','Model','Serial Number','RMA In Tracking'],[['08-03-2026','AP9-A75-NA-R',serial,'']]),
 'Shipping Export':(['Ticket Number','Shipped Date','Model','Serial Number','Tracking Number'],[['M99999999','08-08-2026','AP9-A75-NA-R',serial,'00123456789']]),
 'FedEx Master':(['MSO','Serial Number','Outbound Tracking','Request Date','Fedex Ship Date','Outbound Delivery Date','Return Tracking','Return Tracking Ship Date','Return Tracking Delivery Date'],[['M99999999',serial,'OUT','08-01-2026','08-02-2026','08-03-2026','RETURN','08-04-2026','08-05-2026']])})
r=check(s.post(BASE+'/inventory-management/import',files={'files':('parity.xlsx',p)}));assert '8 events' in r.text,r.text[-1600:]
r=check(s.post(BASE+'/inventory-management/import',files={'files':('parity.xlsx',p)}));assert 'Duplicate file' in r.text
r=check(s.get(BASE+'/inventory-management/serial/'+serial));assert 'fedex_return_delivered' in r.text
r=check(s.get(BASE+'/inventory-management/shipping/all/export.xlsx',params={'preset':'custom','start':'2026-08-01','end':'2026-08-31'}));w=load_workbook(io.BytesIO(r.content));assert any(serial in row for row in w.active.iter_rows(values_only=True))
print('PASS: Training content/attendance/PDF/upload/ZIP/XLSX, Inventory imports/dedup/lifecycle/export')

# The source portal exposes a second Workshop path for pre-sanitized files.
legacy=book({'Repair Data':(['Date Integer','Actual Model','Actual Serial','Derive Size','Type','Type2','Result','Category'],[['08-10-2026','AP9-A75-NA-R','LEGACY-PARITY','75','Depot Repair Tab','Basic','Repaired','Refurbished']])})
r=api('/generate',files={'repair':('legacy.xlsx',legacy),'prev_invoiced':('master.xlsx',master),'shipping':('ship.csv',b'Shipped Date,Serial Number\n')},data={'mode':'legacy','date_from':'2026-08-01','date_to':'2026-08-31','invoice_date':'2026-09-01','completed_date':'2026-08-31','customer':'Promethean','call_id':'C-LEGACY','output_filename':'LEGACY_PARITY'})
d,w=done('/stream');assert d['filename']=='LEGACY_PARITY.xlsx';assert d['subtotal']==110,d;assert w['Breakdown']['E1'].value=='C-LEGACY';assert 'corrected_filename' not in d

# Limited users only see and reach explicitly granted sections.
username='limited_'+uuid.uuid4().hex[:8]
postform('/admin/permissions/users/new',{'username':username,'password':'parity-pass-2026','initials':'LP'})
admin_page=check(s.get(BASE+'/admin/permissions')).text
card=re.search(r'<div class="card"><h2>'+username+r'.*?action="/admin/permissions/users/(\d+)/set"',admin_page,re.S);assert card,username
assert 'name="config_access"' in admin_page and '> Config</label>' in admin_page
uid=card.group(1);postform(f'/admin/permissions/users/{uid}/set',{'invoice_generator':'on','sms_nonconforming':'on','training_role':'viewer'})
limited=requests.Session();check(limited.post(BASE+'/login',data={'username':username,'password':'parity-pass-2026'},allow_redirects=False),302)
portal=check(limited.get(BASE+'/')).text
assert '"trainingRole":"viewer"' in portal and '"config"' not in re.search(r'"invoiceChildren":\[(.*?)\]',portal).group(1)
assert 'id="nav-config"' not in portal and 'id="page-config"' not in portal
check(limited.get(BASE+'/?portal=invoice-generator&client=config'),403)
check(limited.get(BASE+'/get_storage_prices'),403)
viewer_training=check(limited.get(BASE+'/training-tracker/')).text
assert 'portal-nav-admin-permissions' not in viewer_training
assert 'id="portal-nav-my-account"' in viewer_training and '>My Account</span>' in viewer_training
limited_account=check(limited.get(BASE+'/account')).text
assert username in limited_account and 'Standard user' in limited_account
assert 'value="LP" readonly aria-readonly="true"' in limited_account
limited_account=check(limited.post(BASE+'/account/preferences',data={'initials':'LP2','theme_preference':'light'})).text
assert 'Profile and appearance preferences saved.' in limited_account and 'value="light" checked' in limited_account
assert 'value="LP" readonly aria-readonly="true"' in limited_account and 'value="LP2"' not in limited_account
theme_result=check(limited.post(BASE+'/account/theme',data={'theme':'dark'})).json();assert theme_result['ok'] and theme_result['theme']=='dark'
bad_password=check(limited.post(BASE+'/account/password',data={'current_password':'wrong','new_password':'parity-pass-updated','confirm_password':'parity-pass-updated'}),400)
assert 'Current password is incorrect.' in bad_password.text
check(limited.post(BASE+'/account/password',data={'current_password':'parity-pass-2026','new_password':'parity-pass-updated','confirm_password':'parity-pass-updated'}))
check(limited.post(BASE+'/logout',allow_redirects=False),302)
check(limited.post(BASE+'/login',data={'username':username,'password':'parity-pass-updated'},allow_redirects=False),302)
assert 'class="sidebar-sub"' in viewer_training
assert '>Reports</a>' not in viewer_training and '>Content Settings</a>' not in viewer_training
check(limited.post(BASE+'/analyze_amc'),400);check(limited.get(BASE+'/get_philips_dimensions'),403);assert check(limited.get(BASE+'/inventory-management/')).url==BASE+'/'
check(limited.get(BASE+'/training-tracker/'));check(limited.get(BASE+'/training-tracker/admin'),403)
postform(f'/admin/permissions/users/{uid}/delete',{})
print('PASS: legacy Workshop workflow and restricted-user section/role permissions')

# Server-side hardening added by the release audit.
# Reference data can only change through POST with a well-formed body.
check(s.get(BASE+'/set_storage_prices'),405);check(s.get(BASE+'/set_philips_repair_cost'),405);check(s.get(BASE+'/set_amc_prices'),405)
before_amc=api('/get_amc_prices')['prices'];before_tiers=api('/get_philips_repair_cost')['tiers']
r=s.post(BASE+'/set_amc_prices',json={'prices':{'unit_receipt':-1}});check(r,400);assert r.json()['ok'] is False
r=s.post(BASE+'/set_amc_prices',json={'prices':{'made_up_key':5}});check(r,400)
r=s.post(BASE+'/set_philips_repair_cost',json={'tiers':[]});check(r,400)
r=s.post(BASE+'/set_philips_repair_cost',json={'tiers':[{'size':'50','rb_price':'abc','harvest_price':40}]});check(r,400)
r=s.post(BASE+'/set_storage_prices',data='{not json',headers={'Content-Type':'application/json'});check(r,400)
assert api('/get_amc_prices')['prices']==before_amc and api('/get_philips_repair_cost')['tiers']==before_tiers,'rejected writes changed reference data'
r=s.post(BASE+'/set_amc_prices',json={'prices':before_amc});check(r);assert api('/get_amc_prices')['prices']==before_amc
# Cross-site requests carry a foreign Origin and must be refused before any work is done.
r=s.post(BASE+'/set_amc_prices',json={'prices':before_amc},headers={'Origin':'https://evil.example'});check(r,403)
r=s.post(BASE+'/nonconforming/api/items',json={'model':'x','serial':'y','carrier':'z'},headers={'Origin':'https://evil.example'});check(r,403)
r=s.post(BASE+'/set_amc_prices',json={'prices':before_amc},headers={'Origin':ROOT});check(r)
# Logout is a state change and is not triggered by a plain link.
r=s.get(BASE+'/logout',allow_redirects=False);assert r.status_code in (302,303),r.status_code;assert r.headers.get('Location','').rstrip('/').endswith('/index.cfm'),r.headers.get('Location');assert check(s.get(BASE+'/healthz')).json()['ok']
assert 'USSI Nexus' in check(s.get(BASE+'/')).text
# Administrative user creation reports validation problems instead of silently failing.
r=s.post(BASE+'/admin/permissions/users/new',data={'username':'bad user!','password':'parity-pass-2026'},allow_redirects=False);check(r,302);assert 'error=' in r.headers.get('Location','')
r=s.post(BASE+'/admin/permissions/users/new',data={'username':'shortpw_'+uuid.uuid4().hex[:6],'password':'short'},allow_redirects=False);check(r,302);assert 'error=' in r.headers.get('Location','')
assert 'role="alert"' in check(s.get(BASE+'/admin/permissions?error=Example+problem')).text
# Password confirmation is case-sensitive.
bad_confirm=check(s.post(BASE+'/account/password',data={'current_password':TEST_PASSWORD,'new_password':'Parity-Pass-Case-2026','confirm_password':'parity-pass-case-2026'}),400)
print('PASS: POST-only reference data, payload validation, Origin rejection, logout method, admin validation feedback')

# Configurable Philips and TCL rates: defaults reproduce the historical totals, and a
# changed pallet rate flows through the TCL builder without re-analysis.
tclp=api('/get_tcl_prices');assert tclp['prices']['pallet_rate_small']==75 and tclp['defaults']['box_16_20']==15,tclp
phg=api('/get_philips_prices');assert phg['prices']['warehouse_base']==1910 and phg['prices']['inbound_handling']==6,phg
check(s.post(BASE+'/set_tcl_prices',json={'prices':{'pallet_threshold':2.5}}),400)
check(s.post(BASE+'/set_philips_prices',json={'prices':{'warehouse_base':-1}}),400)
api('/set_tcl_prices',json={'prices':{'pallet_rate_small':80}})
api('/confirm_tcl',json={'unit_breakdowns':{d_tcl_key:'1'},'box_breakdowns':{d_tcl_box_key:'1'}})
d,w=done('/stream_tcl');assert d['subtotal']==83.85,d;assert w['Line Items']['G2'].value==80,w['Line Items']['G2'].value
api('/set_tcl_prices',json={'reset':True});assert api('/get_tcl_prices')['prices']['pallet_rate_small']==75

# Workshop corrections: an unusable value is rejected up front and a reviewed exclusion
# is reported in the Excluded Serials tab instead of vanishing.
badraw=book({'Repair Data':(['Date Integer','Actual Model','Actual Serial','Derive Size','Result','Category'],[
 ['08-05-2026','AP9-A75-NA-R','9A75XK101','75','Mainboard replaced','Refurbished'],
 ['08-06-2026','UNKNOWN-MODEL','ZZ99UNRESOLVED','zz','Mainboard replaced','Refurbished']])})
d=api('/sanitize',files={'raw_file':('raw.xlsx',badraw),'prev_invoiced':('master.xlsx',master),'shipping':('ship.csv',ship)},data={'date_from':'2026-08-01','date_to':'2026-08-31','invoice_date':'2026-09-01','completed_date':'2026-08-31','customer':'Promethean','call_id':'C-CORR'})
assert d['issue_count']==1 and d['issues'][0]['issue_type']=='unresolved_size',d
bad_idx=str(d['issues'][0]['row_index'])
r=s.post(BASE+'/generate',data={'output_filename':'CORR_BAD','corrections':json.dumps({bad_idx:{'field':'Derive Size','value':'99'}})});check(r,400);assert 'could not be applied' in r.json()['error'],r.text
api('/generate',data={'output_filename':'CORR_EXCLUDED','corrections':json.dumps({bad_idx:{'field':'Derive Size','value':'EXCLUDE'}})})
d,w=done('/stream');assert d['excluded_count']==1 and d['depot_count']==1,d
excluded_rows=[row for row in w['Excluded Serials'].iter_rows(min_row=2,values_only=True) if any(row)]
assert any('ZZ99UNRESOLVED' in str(row) and 'Excluded by user' in str(row) for row in excluded_rows),excluded_rows
print('PASS: configurable Philips/TCL rates and Workshop correction validation')

# Audit log: actions taken above are visible to the superadmin and exportable.
audit_page=check(s.get(BASE+'/admin/audit',params={'action':'invoice.tcl_built'})).text
assert 'invoice.tcl_built' in audit_page and 'TCL_PARITY.xlsx' in audit_page,audit_page[-2000:]
audit_csv=check(s.get(BASE+'/admin/audit/export.csv',params={'action':'config.tcl_prices'})).text
assert audit_csv.startswith('"Occurred","Actor"') and 'config.tcl_prices' in audit_csv
assert 'auth.login' in check(s.get(BASE+'/admin/audit',params={'actor':TEST_USERNAME})).text
check(limited_audit_probe:=requests.Session().get(BASE+'/admin/audit',allow_redirects=False),302)

# Integration API: keys are created once, scoped, hashed and revocable.
r=s.post(BASE+'/admin/api-keys/new',data={'name':'parity-sync','scopes':'inventory:read,inventory:write,invoices:read,config:read'},allow_redirects=True);check(r)
key_match=re.search(r'nxk_[a-f0-9]{8}_[a-f0-9]{48}',r.text);assert key_match,r.text[-1500:];api_key=key_match.group(0)
assert api_key not in check(s.get(BASE+'/admin/api-keys')).text,'key shown twice'
anon_api=requests.Session()
assert check(anon_api.get(BASE+'/api/v1/health')).json()['api']=='v1'
check(anon_api.get(BASE+'/api/v1/invoices/prices'),401)
check(anon_api.get(BASE+'/api/v1/invoices/prices',headers={'Authorization':'Bearer nxk_deadbeef_'+'0'*48}),401)
H={'Authorization':'Bearer '+api_key}
prices=check(anon_api.get(BASE+'/api/v1/invoices/prices',headers=H)).json();assert prices['ok'] and prices['prices']['tcl']['pallet_rate_small']==75,prices
outs=check(anon_api.get(BASE+'/api/v1/invoices/outputs',headers=H)).json();assert any(o['filename']=='TCL_PARITY.xlsx' for o in outs['outputs']),outs
dl=check(anon_api.get(BASE+'/api/v1/invoices/outputs/TCL_PARITY.xlsx',headers=H));assert dl.content[:2]==b'PK'
ser=check(anon_api.get(BASE+'/api/v1/inventory/serials/'+serial,headers=H)).json();assert ser['serial']['serial_number']==serial,ser
check(anon_api.get(BASE+'/api/v1/inventory/serials/NOPE-'+serial,headers=H),404)
api_serial='API'+uuid.uuid4().hex[:8]
ingest={'kind':'shipping','source':'parity-db','rows':[{'Ticket Number':'M88888888','Shipped Date':'08-09-2026','Model':'AP9-A75-NA-R','Serial Number':api_serial,'Tracking Number':'00987654321'}]}
r=anon_api.post(BASE+'/api/v1/inventory/events',json=ingest,headers=H);check(r,201);assert r.json()['result']['events']>=1,r.text
r=anon_api.post(BASE+'/api/v1/inventory/events',json=ingest,headers=H);check(r,200);assert r.json()['result']['duplicate_payload'] is True
assert check(anon_api.get(BASE+'/api/v1/inventory/serials/'+api_serial,headers=H)).json()['serial']['serial_number']==api_serial
r=anon_api.post(BASE+'/api/v1/inventory/events',json={'kind':'bogus','rows':[{}]},headers=H);check(r,400);assert r.json()['error']['code']=='invalid_body'
r=anon_api.post(BASE+'/api/v1/inventory/events',data='{nope',headers=dict(H,**{'Content-Type':'application/json'}));check(r,400)
ev=check(anon_api.get(BASE+'/api/v1/audit/events',params={'action':'inventory.api_ingest'},headers=H)).json();assert ev['total']>=1 and ev['events'][0]['actor_kind']=='api',ev
keys_page=check(s.get(BASE+'/admin/api-keys')).text;kid=re.search(r'/admin/api-keys/(\d+)/revoke',keys_page).group(1)
postform(f'/admin/api-keys/{kid}/revoke',{});check(anon_api.get(BASE+'/api/v1/invoices/prices',headers=H),401)
print('PASS: audit log pages/export and integration API keys, reads, ingest and revocation')
