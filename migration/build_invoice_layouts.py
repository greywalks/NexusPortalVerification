"""Freeze source workbook presentation, never used by the Lucee runtime.

Usage: python migration/build_invoice_layouts.py ../Logicore-Portal
Reference: greywalks/Logicore-Portal main 02af43e.
Only blank layouts, styles, formulas and headers are retained; no customer rows.
"""
import sys, shutil
from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
SOURCE=Path(sys.argv[1]).resolve();sys.path.insert(0,str(SOURCE))
import amc_builder as amc, philips_builder as philips, storage_builder as storage, builder, tcl_builder as tcl
OUT=Path(__file__).resolve().parents[1]/'template'/'layouts';OUT.mkdir(exist_ok=True)
empty=pd.DataFrame();date=datetime(2000,1,1)
amc.build_amc_invoice(dict(receiving_df=empty,shipping_df=empty,receipt_count=0,ship_count=0,additional_sqft=0),'',OUT/'AMC.xlsx',prices=amc.DEFAULT_PRICES,log=lambda _:None)
philips.build_philips_invoice(dict(demo_additional_sqft=0,service_additional_sqft=0,parts_sqft_manual=0,inbound_count=0,outbound_count=0,repair_count=0,harvest_count=0,repair_total=0,harvest_total=0,received_df=empty,shipping_df=empty,repairs_df=empty),'',OUT/'Philips.xlsx',log=lambda _:None)
storage.build_storage_invoice(dict(unit_storage_df=empty,units_received=empty,programming_df=empty,ship_month=empty,unit_picks_count=0,small_part_picks=0,pallet_count=0,auto_spc_rows=[],unmatched_df=empty),date,date,'','',OUT/'Storage.xlsx',log=lambda _:None)
# A sample row creates the same data-cell styles as a real invoice; erase below.
row={'Actual Model':'','Actual Serial':'','Type2':'Basic','Size':'Small','was_prev_triaged':False,'Unit Price':0,'Result':''}
builder.build(pd.DataFrame([row]),pd.DataFrame([row]),OUT/'Workshop.xlsx',date,date,'','')
for n in range(7):
 path=OUT/f'TCL_{n}.xlsx';shutil.copy(tcl.TEMPLATE_FILE,path);w=load_workbook(path);w._external_links=[]
 tcl._build_invoice_sheet(w,'',date,date,'Contract','Net 30',tcl.DEFAULT_BILL_TO,tcl.DEFAULT_BILL_TO,'',[dict(label='',qty=0,rate=0)]*n,0,0,0)
 tcl._build_line_items_sheet(w,[],[]);w.save(path)
for path in OUT.glob('*.xlsx'):
 w=load_workbook(path);w._external_links=[]
 for ws in w:
  if ws.title in ('Breakdown','Invoice'):continue
  for row in ws.iter_rows(min_row=2):
   for cell in row:cell.value=None
  # Remove stale table dimensions inherited from original sample workbooks.
  for table in ws.tables.values():
   last=table.ref.split(':')[-1];col=''.join(c for c in last if c.isalpha())
   table.ref=f'A1:{col}2'
   if table.autoFilter:table.autoFilter.ref=table.ref
 w.save(path)
print('Created',len(list(OUT.glob('*.xlsx'))),'blank source layouts')
